import copy
import json
import re
from pathlib import Path
from typing import Any

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt

from llm_client import call_llm


def _strip_markdown_fences(text: str) -> str:
    """Remove optional Markdown JSON fences from an LLM response."""
    cleaned = text.strip()
    cleaned = re.sub(r"^```(?:json)?\s*", "", cleaned, flags=re.IGNORECASE)
    cleaned = re.sub(r"\s*```$", "", cleaned)
    return cleaned.strip()


def tailor_resume(resume_json: dict, jd_text: str) -> dict:
    """Tailor a resume against a job description using the configured LLM."""
    if not isinstance(resume_json, dict):
        raise TypeError("resume_json must be a dictionary.")

    if not isinstance(jd_text, str) or not jd_text.strip():
        raise ValueError("jd_text must be a non-empty string.")

    system_prompt = """
You are a resume optimization assistant.

Return ONLY one valid JSON object. Do not use Markdown, code fences,
explanations, commentary, or additional keys.

The JSON object must contain exactly:
{
  "tailored_resume": {},
  "missing_keywords": []
}

Rules:
- tailored_resume must preserve the exact structure of the input resume.
- Do not invent employers, job titles, dates, degrees, certifications,
  technologies, metrics, or accomplishments.
- Reword existing summary and bullets only when supported by the input.
- Include relevant job-description keywords naturally where truthful.
- Reorder skills, experience entries, experience bullets, projects, and
  project bullets by relevance to the job description.
- missing_keywords must contain important job-description terms absent from
  the original resume.
- Keep missing_keywords as a JSON array of concise strings.
""".strip()

    user_prompt = (
        "RESUME JSON:\n"
        f"{json.dumps(resume_json, indent=2, ensure_ascii=True)}\n\n"
        "JOB DESCRIPTION:\n"
        f"{jd_text.strip()}"
    )

    raw_response = call_llm(system_prompt, user_prompt)
    cleaned_response = _strip_markdown_fences(raw_response)

    try:
        result = json.loads(cleaned_response)
    except json.JSONDecodeError as exc:
        preview = cleaned_response[:500]
        raise ValueError(
            "The LLM returned malformed JSON. "
            f"JSON error: {exc}. Response preview: {preview!r}"
        ) from exc

    if not isinstance(result, dict):
        raise ValueError("The LLM response must be a JSON object.")

    if "tailored_resume" not in result or "missing_keywords" not in result:
        raise ValueError(
            'The LLM response must contain "tailored_resume" and '
            '"missing_keywords" keys.'
        )

    if not isinstance(result["tailored_resume"], dict):
        raise ValueError('"tailored_resume" must be a JSON object.')

    if not isinstance(result["missing_keywords"], list):
        raise ValueError('"missing_keywords" must be a JSON array.')

    return result


def _add_bullet(document: Document, text: Any) -> None:
    if text:
        document.add_paragraph(str(text), style="List Bullet")


def export_resume(tailored_resume_json: dict, path: str | Path) -> Path:
    """Export a tailored resume JSON object to a clean DOCX document."""
    if not isinstance(tailored_resume_json, dict):
        raise TypeError("tailored_resume_json must be a dictionary.")

    output_path = Path(path)
    output_path.parent.mkdir(parents=True, exist_ok=True)

    document = Document()
    section = document.sections[0]
    section.top_margin = Inches(0.55)
    section.bottom_margin = Inches(0.55)
    section.left_margin = Inches(0.7)
    section.right_margin = Inches(0.7)

    normal_style = document.styles["Normal"]
    normal_style.font.name = "Calibri"
    normal_style.font.size = Pt(10)

    name = tailored_resume_json.get("name", "Resume")
    title = document.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_run = title.add_run(str(name))
    title_run.bold = True
    title_run.font.size = Pt(18)

    summary = tailored_resume_json.get("summary")
    if summary:
        document.add_heading("Summary", level=1)
        document.add_paragraph(str(summary))

    skills = tailored_resume_json.get("skills", [])
    if skills:
        document.add_heading("Skills", level=1)
        document.add_paragraph(" | ".join(str(skill) for skill in skills))

    experience = tailored_resume_json.get("experience", [])
    if experience:
        document.add_heading("Experience", level=1)
        for item in experience:
            heading = document.add_paragraph()
            heading_run = heading.add_run(
                f"{item.get('title', '')} | {item.get('company', '')}"
            )
            heading_run.bold = True

            dates = " - ".join(
                value
                for value in [
                    item.get("start_date"),
                    item.get("end_date"),
                ]
                if value
            )
            if dates:
                document.add_paragraph(dates)

            for bullet in item.get("bullets", []):
                _add_bullet(document, bullet)

    projects = tailored_resume_json.get("projects", [])
    if projects:
        document.add_heading("Projects", level=1)
        for item in projects:
            heading = document.add_paragraph()
            heading_run = heading.add_run(str(item.get("name", "Project")))
            heading_run.bold = True

            if item.get("description"):
                document.add_paragraph(str(item["description"]))

            technologies = item.get("technologies", [])
            if technologies:
                document.add_paragraph(
                    "Technologies: "
                    + ", ".join(str(value) for value in technologies)
                )

            for bullet in item.get("bullets", []):
                _add_bullet(document, bullet)

    education = tailored_resume_json.get("education", [])
    if education:
        document.add_heading("Education", level=1)
        for item in education:
            line = " | ".join(
                str(value)
                for value in [
                    item.get("degree"),
                    item.get("institute"),
                    item.get("year"),
                ]
                if value
            )
            if line:
                document.add_paragraph(line)

    document.save(output_path)
    return output_path